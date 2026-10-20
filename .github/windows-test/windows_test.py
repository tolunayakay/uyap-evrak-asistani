# -*- coding: utf-8 -*-
"""
Windows doğrulama testi (GitHub Actions'ta çalışır).

Aynı test seti hem orijinal kodla (main) hem macOS uyumluluk dalıyla çalıştırılır; sonuçlar
karşılaştırılarak Windows davranışının değişmediği doğrulanır. Dosya pencereleri ve uyarılar
sunucuda kullanıcı olmadığı için kodla yanıtlanır; işlemler butonların kullandığı yoldan
(_thread_calistir) çalıştırılır.

Kullanım: python windows_test.py <uyap_asistan.py'nin bulunduğu klasör> <çıktı klasörü>
"""
import io
import json
import os
import shutil
import sys
import threading
import time
import traceback
import zipfile

KAYNAK = os.path.abspath(sys.argv[1])
CIKTI = os.path.abspath(sys.argv[2])
T = os.path.join(CIKTI, "dosyalar")
shutil.rmtree(CIKTI, ignore_errors=True)
os.makedirs(T)
os.chdir(KAYNAK)
sys.path.insert(0, KAYNAK)
sys.argv = ["uyap_asistan.py"]

import docx
import pymupdf
from PIL import Image, ImageGrab
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog
import customtkinter as ctk

import uyap_asistan as u

# ---------------------------------------------------------------- test verisi
import html


def pdf_yap(yol, sayfalar):
    # PyMuPDF'in yerleşik yazı tipi: Türkçe karakterler ve normal boşluklar her sistemde aynı geri okunur
    d = pymupdf.open()
    for metin in sayfalar:
        p = d.new_page()
        if metin:
            p.insert_htmlbox(pymupdf.Rect(60, 60, 540, 780),
                             '<p style="line-height:2">' + html.escape(metin).replace("\n", "<br>") + "</p>")
    d.set_metadata({"author": "Gizli Yazar", "creator": "BilgisayarAdi-PC", "title": "Gizli Başlık"})
    d.save(yol)


METIN = ("T.C. İSTANBUL 3. ASLİYE HUKUK MAHKEMESİ\nESAS NO: 2026/123\nKARAR NO: 2026/456\nDAVACI: Ahmet Yılmaz\n"
         "DAVALI: Ayşe Şahin\nT.C. Kimlik No: 10000000146\nIBAN: TR330006100519786457841326\nTelefon: 0532 123 45 67\n"
         "KONU: Kira bedelinin tespiti talebidir. Davacı, davalının kiracı olarak oturduğu taşınmazın kira bedelinin "
         "güncel rayiç değere göre belirlenmesini talep etmiştir. Mahkemece bilirkişi incelemesi yaptırılmıştır.\n"
         "HÜKÜM: Davanın KABULÜ ile aylık kira bedelinin 25.000 TL olarak tespitine karar verilmiştir.")
pdf_yap(f"{T}/dava.pdf", [METIN, "", "İkinci içerik sayfası çğıöşü ÇĞİÖŞÜ", ""])
pdf_yap(f"{T}/ek2.pdf", ["Ek belge sayfası 1", "Ek belge sayfası 2"])
Image.new("RGB", (1200, 1600), (200, 220, 255)).save(f"{T}/tarama.jpg", quality=90)
Image.new("RGBA", (800, 600), (255, 0, 0, 128)).save(f"{T}/seffaf.png")
import random
random.seed(1)
d = pymupdf.open()
for i in range(6):
    p = d.new_page()
    img = Image.frombytes("RGB", (1400, 1400), bytes(random.getrandbits(8) for _ in range(1400 * 1400 * 3)))
    b = io.BytesIO(); img.save(b, "PNG"); p.insert_image(p.rect, stream=b.getvalue())
d.save(f"{T}/buyuk.pdf")
w = docx.Document()
w.add_heading("DİLEKÇE", 1)
w.add_paragraph("Sayın Mahkeme, müvekkilimiz adına işbu dilekçeyi sunuyoruz. Türkçe: çğıöşü ÇĞİÖŞÜ")
t = w.add_table(rows=2, cols=2)
t.cell(0, 0).text = "Taraf"; t.cell(0, 1).text = "Ad"; t.cell(1, 0).text = "Davacı"; t.cell(1, 1).text = "Ali Veli"
w.add_picture(f"{T}/tarama.jpg", width=docx.shared.Cm(4))
w.core_properties.author = "Gizli Yazar"; w.core_properties.last_modified_by = "Gizli Düzenleyen"
w.save(f"{T}/dilekce.docx")

# ---------------------------------------------------------------- pencere yanıtları
kayit = []          # (tür, başlık, mesaj)
cevaplar = []       # sıradaki dosya penceresi / değer girişi cevapları


def sonraki(varsayilan=""):
    return cevaplar.pop(0) if cevaplar else varsayilan


for _ad in ("showinfo", "showwarning", "showerror"):
    setattr(messagebox, _ad, (lambda ad: lambda baslik="", mesaj="", **k: kayit.append((ad, baslik, mesaj)))(_ad))
messagebox.askyesno = lambda baslik="", mesaj="", **k: (kayit.append(("askyesno", baslik, mesaj)), True)[1]
filedialog.asksaveasfilename = lambda **k: sonraki()
filedialog.askopenfilename = lambda **k: sonraki()
filedialog.askopenfilenames = lambda **k: sonraki(())
filedialog.askdirectory = lambda **k: sonraki()
simpledialog.askstring = lambda *a, **k: sonraki()
acilanlar = []
os.startfile = lambda yol, *a: acilanlar.append(yol)       # Windows'ta dosya açma
if sys.platform != "win32" and hasattr(u, "dosyayi_ac"):    # Mac'te kuru deneme için
    u.dosyayi_ac = lambda yol: acilanlar.append(yol)

app = u.ModernUyapAsistani()
app.geometry("1180x760+20+20")
app.update()


def pompa(sn=0.3):
    bitis = time.time() + sn
    while time.time() < bitis:
        app.update(); time.sleep(0.02)


def bekle(azami=240):
    bitis = time.time() + azami
    pompa(0.3)
    while time.time() < bitis:
        isler = [t for t in threading.enumerate() if t.is_alive() and t is not threading.main_thread()
                 and t.name.startswith("Thread")]     # _thread_calistir'ın başlattığı iş parçacıkları
        if not isler:
            break
        app.update(); time.sleep(0.05)
    pompa(0.3)


def pencereler():
    return [w for w in app.winfo_children() if isinstance(w, (tk.Toplevel, ctk.CTkToplevel))]


def tum_cocuklar(w):
    for c in w.winfo_children():
        yield c; yield from tum_cocuklar(c)


def tikla(parca, kok=None):
    for kokw in ([kok] if kok else pencereler()[::-1]):
        for c in tum_cocuklar(kokw):
            if isinstance(c, ctk.CTkButton) and parca in str(c.cget("text")):
                c.invoke(); return True
    raise RuntimeError(f"Buton bulunamadı: {parca}")


def kapat_hepsi():
    for p in pencereler():
        try: p.destroy()
        except Exception: pass
    pompa(0.2)


def sec(yol):
    if yol not in app.dosyalar:
        app.dosya_ekle_listeye(yol)
    app.secili_dosya_indeksi = app.dosyalar.index(yol)
    app._liste_arayuzunu_yenile(); pompa(0.1)


def buton(fn):
    """İşlemi butonun yaptığı gibi arka plan iş parçacığında başlatır ve bitmesini bekler."""
    app._thread_calistir(fn); bekle()


def sayfa_say(yol):
    with pymupdf.open(yol) as d:
        return len(d)


def metin(yol):
    with pymupdf.open(yol) as d:
        return "".join(p.get_text() for p in d).replace("\xa0", " ")


def ekran(ad):
    try:
        app.update(); app.lift(); app.attributes("-topmost", True); pompa(0.6)
        x, y, w_, h = app.winfo_rootx(), app.winfo_rooty(), app.winfo_width(), app.winfo_height()
        ImageGrab.grab(bbox=(x, y, x + w_, y + h)).save(os.path.join(CIKTI, f"ekran_{ad}.png"))
        app.attributes("-topmost", False)
    except Exception as e:
        print("ekran görüntüsü alınamadı:", e)


sonuclar = []
TESTLER = []


def test(ad):
    def kaydet(fn):
        TESTLER.append((ad, fn))
    return kaydet


def calistir(ad, fn):
    kayit.clear(); acilanlar.clear(); cevaplar.clear()
    t0 = time.time()
    try:
        notu = fn()
        hatalar = [k for k in kayit if k[0] == "showerror"]
        if hatalar and not ad.startswith("Word ➔ PDF"):
            raise RuntimeError(f"Uygulama hata gösterdi: {hatalar[0][1]}: {hatalar[0][2][:200]}")
        sonuclar.append({"test": ad, "durum": "GEÇTİ", "not": str(notu or ""), "sure": round(time.time() - t0, 1)})
    except Exception as e:
        sonuclar.append({"test": ad, "durum": "KALDI", "not": f"{type(e).__name__}: {str(e)[:200]}",
                         "sure": round(time.time() - t0, 1)})
        traceback.print_exc()
    kapat_hepsi()
    print(f"[{sonuclar[-1]['durum']}] {ad} — {sonuclar[-1]['not']}", flush=True)


# ---------------------------------------------------------------- testler
@test("Platform ayarları (yazı tipi, ayar klasörü, Türkçe yazı tipi)")
def _():
    arayuz = getattr(u, "ARAYUZ_FONTU", "Segoe UI")
    assert arayuz == "Segoe UI", arayuz
    assert getattr(u, "IS_MAC", False) is False
    assert u.AYAR_DOSYA_YOLU.lower().startswith(os.environ["APPDATA"].lower()), u.AYAR_DOSYA_YOLU
    assert u.TURKCE_FONT_YOLU and u.TURKCE_FONT_YOLU.lower().endswith("arial.ttf"), u.TURKCE_FONT_YOLU
    return f"yazı tipi={arayuz} | ayarlar={u.AYAR_DOSYA_YOLU} | pdf yazı tipi={u.TURKCE_FONT_YOLU}"


@test("Ana pencere ve ekran görüntüsü (açık tema)")
def _():
    ekran("acik_tema")
    return "pencere açık"


@test("Dosya listesine ekleme (PDF, JPG, PNG, DOCX, klasör)")
def _():
    os.makedirs(f"{T}/klasor", exist_ok=True); shutil.copy(f"{T}/ek2.pdf", f"{T}/klasor/klasordeki.pdf")
    for y in ("dava.pdf", "ek2.pdf", "tarama.jpg", "seffaf.png", "dilekce.docx", "klasor"):
        app.dosya_ekle_listeye(f"{T}/{y}")
    assert len(app.dosyalar) == 6, app.dosyalar
    return f"{len(app.dosyalar)} dosya"


@test("Sürükle-bırak olayı (boşluklu/Türkçe adlı dosya)")
def _():
    yol = f"{T}/Boşluklu Dosya Adı İğüş.pdf"; shutil.copy(f"{T}/ek2.pdf", yol)

    class Olay:
        data = app.tk.call("list", yol)
    app.dosyalari_surukle_birak(Olay())
    assert yol in app.dosyalar
    app.dosyalar.remove(yol); app._liste_arayuzunu_yenile()
    return "tkdnd " + str(getattr(app, "TkdndVersion", "?"))


@test("Sekme geçişleri")
def _():
    tv = app.tabview
    komut = tv._segmented_button._command
    sira = ["Araçlar", "Birleştir", "Dönüştür", "Sayfalar", "Birleştir", "Araçlar"]
    for ad in sira:
        komut(ad); pompa(0.4)
        gorunen = [n for n, f in tv._tab_dict.items() if f.winfo_ismapped()]
        assert tv.get() == ad and gorunen == [ad], (ad, tv.get(), gorunen)
    return " → ".join(sira)


@test("Birleştirme + Bates + Ek Listesi (fihrist)")
def _():
    app.dosyalar[:] = [f"{T}/dava.pdf", f"{T}/ek2.pdf", f"{T}/tarama.jpg", f"{T}/seffaf.png"]
    app._liste_arayuzunu_yenile()
    app.bates_aktif.set(True); app.fihrist_aktif.set(True)
    cevaplar[:] = [f"{T}/birlesik.pdf"]
    buton(app.tumunu_birlestir_pdf_yap)
    n = sayfa_say(f"{T}/birlesik.pdf"); m = metin(f"{T}/birlesik.pdf")
    with pymupdf.open(f"{T}/birlesik.pdf") as d:
        toc = d.get_toc()
    assert n == 9 and "EK LİSTESİ" in m and len(toc) >= 4, (n, len(toc))
    app.bates_aktif.set(False); app.fihrist_aktif.set(False)
    return f"{n} sayfa, {len(toc)} yer imi"


@test("Sıkıştırma (Akıllı mod)")
def _():
    sec(f"{T}/buyuk.pdf"); cevaplar[:] = [f"{T}/sikisik.pdf"]
    buton(app.pdf_tekil_sikistir)
    assert sayfa_say(f"{T}/sikisik.pdf") == 6
    return f"{os.path.getsize(f'{T}/sikisik.pdf') / 1e6:.1f} MB"


@test("Sıkıştırma (Agresif raster 150 DPI)")
def _():
    eski = app.mod_secimi.get()
    for c in tum_cocuklar(app):
        if isinstance(c, ctk.CTkRadioButton) and "Agresif" in str(c.cget("text")):
            c.invoke()
    sec(f"{T}/buyuk.pdf"); cevaplar[:] = [f"{T}/agresif.pdf"]
    buton(app.pdf_tekil_sikistir)
    app.mod_secimi.set(eski)
    assert sayfa_say(f"{T}/agresif.pdf") == 6
    return f"{os.path.getsize(f'{T}/agresif.pdf') / 1e6:.1f} MB"


@test("UYAP MB bölücü (değer girişi + klasör)")
def _():
    sec(f"{T}/buyuk.pdf"); os.makedirs(f"{T}/parcalar")
    sinir = max(0.5, os.path.getsize(f"{T}/buyuk.pdf") / 1e6 / 3)
    cevaplar[:] = [f"{sinir:.1f}".replace(".", ","), f"{T}/parcalar"]
    buton(app.pdf_akilli_boyut_bol)
    parcalar = sorted(os.listdir(f"{T}/parcalar"))
    assert len(parcalar) >= 2 and sum(sayfa_say(f"{T}/parcalar/{p}") for p in parcalar) == 6, parcalar
    return f"{len(parcalar)} parça"


@test("Sayfa döndürme")
def _():
    sec(f"{T}/dava.pdf"); app.pdf_dondur_penceresi(); pompa()
    for c in tum_cocuklar(pencereler()[-1]):
        if isinstance(c, ctk.CTkEntry):
            c.insert(0, "1-2")
    cevaplar[:] = [f"{T}/donmus.pdf"]; tikla("Uygula ve Kaydet"); bekle()
    with pymupdf.open(f"{T}/donmus.pdf") as d:
        r = [p.rotation for p in d]
    assert r == [90, 90, 0, 0], r
    return str(r)


@test("Boş sayfa temizleme")
def _():
    sec(f"{T}/dava.pdf"); cevaplar[:] = [f"{T}/bossuz.pdf"]
    buton(app.pdf_bos_sayfalari_temizle)
    assert sayfa_say(f"{T}/bossuz.pdf") == 2
    return "4 → 2 sayfa"


@test("Sayfa sayfa bölme")
def _():
    sec(f"{T}/dava.pdf"); os.makedirs(f"{T}/tekil"); cevaplar[:] = [f"{T}/tekil"]
    buton(app.pdf_sayfa_sayfa_bol)
    assert len(os.listdir(f"{T}/tekil")) == 4
    return "4 dosya"


@test("Aralık çıkarma (1, 3)")
def _():
    sec(f"{T}/dava.pdf"); cevaplar[:] = ["1, 3", f"{T}/aralik.pdf"]
    app.pdf_aralik_bol(); bekle()
    assert sayfa_say(f"{T}/aralik.pdf") == 2 and "İkinci içerik" in metin(f"{T}/aralik.pdf")
    return "2 sayfa"


@test("Görsel sayfa önizleme penceresi")
def _():
    sec(f"{T}/dava.pdf"); app.pdf_gorsel_onizleme_ac(); pompa(2)
    assert [w for w in pencereler() if isinstance(w, u.PdfGorselOnizlemePenceresi)]
    return "açıldı"


@test("PDF ➔ UDF (dosyayı otomatik açma dahil)")
def _():
    sec(f"{T}/dava.pdf"); app.pdf_to_udf_dialog(); pompa()
    cevaplar[:] = [f"{T}/dava.udf"]; tikla("UDF'e Dönüştür"); bekle()
    with zipfile.ZipFile(f"{T}/dava.udf") as z:
        x = z.read("content.xml").decode("utf-8")
    assert "Ahmet Yılmaz" in x and "KABULÜ" in x
    assert acilanlar and acilanlar[-1].endswith("dava.udf"), acilanlar   # os.startfile çağrıldı mı
    return "os.startfile çağrıldı"


@test("PDF ➔ Word")
def _():
    sec(f"{T}/dava.pdf"); cevaplar[:] = [f"{T}/dava.docx"]
    app.pdf_to_docx_dialog(); bekle()
    m = "\n".join(p.text for p in docx.Document(f"{T}/dava.docx").paragraphs)
    assert "Ayşe Şahin" in m
    return "metin doğru"


@test("Word ➔ UDF")
def _():
    sec(f"{T}/dilekce.docx"); cevaplar[:] = [f"{T}/dilekce.udf"]
    buton(app.word_to_udf)
    with zipfile.ZipFile(f"{T}/dilekce.udf") as z:
        x = z.read("content.xml").decode("utf-8")
    assert "çğıöşü" in x and "Ali Veli" in x
    return "tablo/resim: " + str(x.count("<table") + x.count("<image"))


@test("UDF ➔ Word")
def _():
    sec(f"{T}/dilekce.udf"); cevaplar[:] = [f"{T}/udfden.docx"]
    buton(app.udf_to_word)
    d = docx.Document(f"{T}/udfden.docx")
    assert "çğıöşü" in "\n".join(p.text for p in d.paragraphs)
    return f"{len(d.tables)} tablo, {len(d.inline_shapes)} resim"


@test("Metin önizleme (PDF / Word / UDF) + pano")
def _():
    for y in ("dava.pdf", "dilekce.docx", "dilekce.udf"):
        sec(f"{T}/{y}"); app.metin_oku_penceresi(); pompa()
        kutu = [c for c in tum_cocuklar(pencereler()[-1]) if isinstance(c, ctk.CTkTextbox)][0]
        assert "Ahmet" in kutu.get("0.0", "end") or "Sayın Mahkeme" in kutu.get("0.0", "end"), y
        tikla("Panoya Kopyala"); pompa()
        assert len(app.clipboard_get()) > 20
        kapat_hepsi()
    return "3 biçim"


@test("e-İmza ön kontrolü (imzasız belge)")
def _():
    sec(f"{T}/dava.pdf"); app.eimza_dogrula_penceresi(); pompa()
    assert pencereler()
    return "pencere açıldı"


@test("Aslı Gibidir kaşesi")
def _():
    sec(f"{T}/dava.pdf"); app.pdf_kase_ekle_penceresi(); pompa()
    girdiler = [c for c in tum_cocuklar(pencereler()[-1]) if isinstance(c, ctk.CTkEntry)]
    girdiler[1].delete(0, "end"); girdiler[1].insert(0, "Av. Deneme Çağrı Öztürk")
    cevaplar[:] = [f"{T}/kaseli.pdf"]; tikla("Kaşeyi Bas"); bekle()
    m = metin(f"{T}/kaseli.pdf")
    assert "ASLI GİBİDİR" in m and "Öztürk" in m
    return "kaşe metni PDF'te"


@test("Filigran")
def _():
    sec(f"{T}/dava.pdf"); app.pdf_filigran_ekle_penceresi(); pompa()
    tikla("TASLAKTIR"); pompa()
    cevaplar[:] = [f"{T}/filigranli.pdf"]; tikla("Filigranı Bas"); bekle()
    assert "TASLAKTIR" in metin(f"{T}/filigranli.pdf")
    return "filigran PDF'te"


@test("KVKK karartma (TC, IBAN, telefon, özel kelime)")
def _():
    sec(f"{T}/dava.pdf"); app.akilli_karartma_penceresi(); pompa()
    for c in tum_cocuklar(pencereler()[-1]):
        if isinstance(c, ctk.CTkEntry):
            c.insert(0, "Ahmet Yılmaz")
        if isinstance(c, ctk.CTkCheckBox) and "Telefon" in str(c.cget("text")) and not c.get():
            c.toggle()
    cevaplar[:] = [f"{T}/karartilmis.pdf"]; tikla("Kalıcı Karart"); bekle()
    m = metin(f"{T}/karartilmis.pdf")
    for s in ("10000000146", "TR330006100519786457841326", "0532 123 45 67", "Ahmet Yılmaz"):
        assert s not in m, s
    assert "Ayşe Şahin" in m
    return "4 tür veri silindi"


@test("Metaveri kazıma (PDF + Word)")
def _():
    app.dosyalar[:] = [f"{T}/dava.pdf", f"{T}/dilekce.docx"]; app.secili_dosya_indeksi = 0; app._liste_arayuzunu_yenile()
    app.metaveri_kazima_penceresi(); pompa()
    for c in tum_cocuklar(pencereler()[-1]):
        if isinstance(c, ctk.CTkRadioButton) and "tüm PDF" in str(c.cget("text")):
            c.invoke()
    os.makedirs(f"{T}/temiz"); cevaplar[:] = [f"{T}/temiz"]
    tikla("Metaverileri Kazı"); bekle()
    with pymupdf.open(f"{T}/temiz/temiz_dava.pdf") as d:
        md = d.metadata
    cp = docx.Document(f"{T}/temiz/temiz_dilekce.docx").core_properties
    assert "Gizli" not in str(md) and "Gizli" not in (cp.author or "") + (cp.last_modified_by or "")
    return "temiz"


@test("Word ➔ PDF (Word yok: hata mesajı, uygulama açık kalmalı)")
def _():
    sec(f"{T}/dilekce.docx"); cevaplar[:] = [f"{T}/wordden.pdf"]
    buton(app.word_to_pdf)
    hatalar = [k for k in kayit if k[0] == "showerror"]
    assert app.winfo_exists()
    if os.path.exists(f"{T}/wordden.pdf"):
        return "Word bulundu ve PDF üretildi"
    assert hatalar, kayit
    return "hata mesajı: " + hatalar[0][2][:120].replace("\n", " ")


@test("Ayarlar penceresi + tema + koyu tema ekran görüntüsü")
def _():
    app.ayarlar_penceresi_ac(); pompa(1)
    assert pencereler()
    kapat_hepsi()
    once = ctk.get_appearance_mode()
    app.tema_switch.toggle(); pompa(1)
    sonra = ctk.get_appearance_mode()
    assert sonra != once, (once, sonra)
    ekran("koyu_tema" if sonra == "Dark" else "acik_tema_2")
    app.tema_switch.toggle(); pompa(0.5)
    assert ctk.get_appearance_mode() == once
    return f"{once} → {sonra} → {once}"


@test("Nasıl Kullanılır rehberi")
def _():
    app.rehber_penceresi_ac(); pompa(1)
    assert pencereler()
    return "açıldı"


def hepsini_calistir():
    for ad, fn in TESTLER:
        calistir(ad, fn)
    app.quit()


app.after(300, hepsini_calistir)
app.mainloop()
try:
    app.destroy()
except Exception:
    pass

with open(os.path.join(CIKTI, "sonuclar.json"), "w", encoding="utf-8") as f:
    json.dump(sonuclar, f, ensure_ascii=False, indent=1)
gecen = sum(s["durum"] == "GEÇTİ" for s in sonuclar)
print(f"\n{gecen}/{len(sonuclar)} geçti")
